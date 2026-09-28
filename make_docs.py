import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import json
import docx2pdf

print('[1/4] Loading metadata...')
with open('docs/frontend_data.json', 'r', encoding='utf-8') as f:
    frontend_files = json.load(f)

with open('docs/backend_data.json', 'r', encoding='utf-8') as f:
    backend_files = json.load(f)

print(f'Loaded {len(frontend_files)} frontend files and {len(backend_files)} backend files.')

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_rtl_p(doc, text='', align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=2, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    
    if text:
        add_rtl_text(p, text)
    return p

def add_rtl_text(p, text, font_name='Arial', size_pt=10.5, bold=False, italic=False, color_rgb=(40, 40, 40)):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rtl {nsdecls("w")}/>'))
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:cs="{font_name}" w:ascii="{font_name}" w:hAnsi="{font_name}"/>')
    rPr.append(rFonts)
    return run

def add_h1(doc, text):
    p = add_rtl_p(doc, space_before=18, space_after=8)
    add_rtl_text(p, text, font_name='Arial', size_pt=16, bold=True, color_rgb=(30, 58, 138))
    return p

def add_file_card(doc, file_data):
    file_name = file_data['file_name']
    file_path = file_data['file_path']
    category_tag = file_data['category']
    role_desc = file_data['role']
    key_items = file_data['key_items']
    doctor_q = file_data['doctor_q']
    doctor_a = file_data['doctor_a']

    table = doc.add_table(rows=5, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    c0 = table.rows[0].cells[0]
    c0.width = Inches(7.0)
    set_cell_background(c0, '1E3A8A')
    set_cell_margins(c0, top=110, bottom=110, left=150, right=150)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_rtl_text(p0, f'📄 {file_name}  [{category_tag}]', font_name='Arial', size_pt=11.5, bold=True, color_rgb=(255, 255, 255))
    
    c1 = table.rows[1].cells[0]
    c1.width = Inches(7.0)
    set_cell_background(c1, 'F1F5F9')
    set_cell_margins(c1, top=70, bottom=70, left=150, right=150)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_rtl_text(p1, 'المسار البرمجي: ', font_name='Arial', size_pt=9.5, bold=True, color_rgb=(71, 85, 105))
    add_rtl_text(p1, file_path, font_name='Consolas', size_pt=9.5, bold=False, color_rgb=(15, 23, 42))

    c2 = table.rows[2].cells[0]
    c2.width = Inches(7.0)
    set_cell_background(c2, 'FFFFFF')
    set_cell_margins(c2, top=90, bottom=90, left=150, right=150)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_rtl_text(p2, '📌 الدور الوظيفي والمعماري: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(30, 41, 59))
    add_rtl_text(p2, role_desc, font_name='Arial', size_pt=10, color_rgb=(51, 65, 85))

    c3 = table.rows[3].cells[0]
    c3.width = Inches(7.0)
    set_cell_background(c3, 'F8FAFC')
    set_cell_margins(c3, top=90, bottom=90, left=150, right=150)
    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p3._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_rtl_text(p3, '🔑 أهم الكلاسات والدوال والمكونات: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(15, 118, 110))
    add_rtl_text(p3, key_items, font_name='Arial', size_pt=10, color_rgb=(51, 65, 85))

    c4 = table.rows[4].cells[0]
    c4.width = Inches(7.0)
    set_cell_background(c4, 'EFF6FF')
    set_cell_margins(c4, top=110, bottom=110, left=150, right=150)
    p4 = c4.paragraphs[0]
    p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p4._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    add_rtl_text(p4, '❓ سؤال الدكتور المتوقع: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(180, 83, 9))
    add_rtl_text(p4, f'"{doctor_q}"\n', font_name='Arial', size_pt=10, italic=True, bold=True, color_rgb=(146, 64, 14))
    add_rtl_text(p4, '💡 الإجابة النموذجية: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(21, 128, 61))
    add_rtl_text(p4, doctor_a, font_name='Arial', size_pt=10, color_rgb=(22, 101, 52))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

print('[2/4] Writing Header and Overview...')
p_title = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=4)
add_rtl_text(p_title, 'تطبيق تاسك فلو (TaskFlow)', font_name='Arial', size_pt=24, bold=True, color_rgb=(30, 58, 138))

p_sub = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=14)
add_rtl_text(p_sub, 'الدليل المرجعي الهندسي الشامل لجميع ملفات النظام (Frontend lib و Backend FastAPI)', font_name='Arial', size_pt=12, bold=True, color_rgb=(71, 85, 105))

p_meta = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
add_rtl_text(p_meta, 'المطور: م/ الحاج  |  المشروع: نظام متكامل لإدارة المهام والإنتاجية  |  سبتمبر 2026', font_name='Arial', size_pt=10, italic=True, color_rgb=(100, 116, 139))

p_intro = add_rtl_p(doc, space_before=4, space_after=14)
add_rtl_text(p_intro, 'الملخص التنفيذي للمعمارية: ', font_name='Arial', size_pt=11, bold=True, color_rgb=(30, 41, 59))
add_rtl_text(p_intro, 'تم تصميم وبناء تطبيق TaskFlow وفق معمارية Clean Architecture المتقدمة مع دعم كامل لمبدأ Offline-First. يحتوي النظام على واجهة مستخدم حديثة مبنية بإطار عمل Flutter بلغة Dart 3 ومكتبة Provider لإدارة الحالة وقاعدة بيانات SQLite المحلية، متصلة عبر شبكة آمنة مع خادم خلفي عالي الأداء مبني بواسطة FastAPI و SQLAlchemy 2.0 ومحرك PostgreSQL. يوثق هذا الدليل كافة ملفات المشروع الـ 54 ملفاً، موضحاً دور كل ملف وأهم مكوناته وسؤال المناقشة المتوقع وإجابته النموذجية.', font_name='Arial', size_pt=10.5, color_rgb=(51, 65, 85))

print('[3/4] Adding Frontend and Backend Sections...')
add_h1(doc, 'الجزء الأول: ملفات الواجهة الأمامية وتطبيق الهاتف (Flutter - lib)')
p_f_info = add_rtl_p(doc, space_before=2, space_after=10)
add_rtl_text(p_f_info, f'يحتوي مجلد lib على {len(frontend_files)} ملفاً برمجياً موزعاً على طبقات النماذج، المزودات، الخدمات، الشاشات، والمكونات المشتركة.', font_name='Arial', size_pt=10.5, italic=True, color_rgb=(71, 85, 105))

for f_data in frontend_files:
    add_file_card(doc, f_data)

doc.add_page_break()
add_h1(doc, 'الجزء الثاني: ملفات الخادم الخلفي وقواعد البيانات (Backend - FastAPI & PostgreSQL)')
p_b_info = add_rtl_p(doc, space_before=2, space_after=10)
add_rtl_text(p_b_info, f'يحتوي مجلد Backend على {len(backend_files)} ملفاً برمجياً موزعاً على طبقات النواة، اتصال قاعدة البيانات، نماذج ORM، مخططات التحقق Pydantic، مسارات الـ API، والخدمات الأمنية.', font_name='Arial', size_pt=10.5, italic=True, color_rgb=(71, 85, 105))

for b_data in backend_files:
    add_file_card(doc, b_data)

doc_path = 'docs/TaskFlow_Architecture_Manual.docx'
doc.save(doc_path)
print(f'DOCX successfully created at {doc_path}')

print('[4/4] Converting DOCX to PDF...')
pdf_path = 'docs/TaskFlow_Architecture_Manual.pdf'
docx2pdf.convert(doc_path, pdf_path)
print(f'SUCCESS: PDF created at {pdf_path}')
