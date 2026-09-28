import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import json
import docx2pdf

print('[1/4] Initializing Lib Code Manual Builder...')

with open('docs/frontend_data.json', 'r', encoding='utf-8') as f:
    frontend_meta = json.load(f)

meta_lookup = {item['file_path'].replace('\\', '/'): item for item in frontend_meta}

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_rtl_p(doc, text='', align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=2, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
    
    if text:
        add_rtl_text(p, text)
    return p

def add_rtl_text(p, text, font_name='Arial', size_pt=10.5, bold=False, italic=False, color_rgb=(40, 40, 40)):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    
    rPr = run._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rtl {nsdecls("w")}/>'))
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:cs="{font_name}" w:ascii="{font_name}" w:hAnsi="{font_name}"/>')
    rPr.append(rFonts)
    return run

def add_h1(doc, text):
    p = add_rtl_p(doc, space_before=18, space_after=8)
    add_rtl_text(p, text, font_name='Arial', size_pt=16, bold=True, color_rgb=(30, 58, 138))
    return p

def add_h2(doc, text):
    p = add_rtl_p(doc, space_before=12, space_after=6)
    add_rtl_text(p, text, font_name='Arial', size_pt=13, bold=True, color_rgb=(15, 118, 110))
    return p

def add_code_box(doc, code_str, max_lines=180):
    lines = code_str.split('\n')
    if len(lines) > max_lines:
        display_lines = lines[:max_lines]
        truncated = True
    else:
        display_lines = lines
        truncated = False

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(7.0)
    set_cell_background(cell, '0F172A') # Dark Slate / VS Code Dark
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    
    formatted_code = '\n'.join(display_lines)
    if truncated:
        formatted_code += f'\n\n// ... [تم اختصار بقية الكود البالغ {len(lines)} سطراً للحفاظ على انسيابية الملف] ...'
        
    run = p.add_run(formatted_code)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(226, 232, 240) # Slate light
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(8)

# Header & Overview
p_title = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=4)
add_rtl_text(p_title, 'الأكواد المصدرية الكاملة لتطبيق تاسك فلو (TaskFlow - lib)', font_name='Arial', size_pt=22, bold=True, color_rgb=(30, 58, 138))

p_sub = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=14)
add_rtl_text(p_sub, 'مرجع شامل: الكود المصدري الأصلي لكل ملف مع الشرح الهندسي وسؤال المناقشة المتوقع', font_name='Arial', size_pt=12, bold=True, color_rgb=(71, 85, 105))

p_meta = add_rtl_p(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
add_rtl_text(p_meta, 'المطور: م/ الحاج  |  إطار العمل: Flutter (Dart 3)  |  التاريخ: سبتمبر 2026', font_name='Arial', size_pt=10, italic=True, color_rgb=(100, 116, 139))

p_intro = add_rtl_p(doc, space_before=4, space_after=14)
add_rtl_text(p_intro, 'مقدمة توثيق الكود: ', font_name='Arial', size_pt=11, bold=True, color_rgb=(30, 41, 59))
add_rtl_text(p_intro, 'يحتوي هذا المرجع المتخصص على الشيفرة المصدرية (Source Code) لجميع ملفات الواجهة الأمامية والخدمات وإدارة الحالة داخل مجلد lib البالغ عددها 29 ملفاً. تم تذييل كل ملف بشرح تحليلي لدوره البرمجي والمعماري، وتفصيل أهم مكوناته، مع استعراض سؤال المناقشة المتوقع من الدكتور وإجابته النموذجية.', font_name='Arial', size_pt=10.5, color_rgb=(51, 65, 85))

print('[2/4] Reading Dart Files and Building Documentation...')

file_count = 0
for root, dirs, files in os.walk('lib'):
    for file in sorted(files):
        if file.endswith('.dart'):
            file_count += 1
            full_path = os.path.join(root, file)
            rel_path = os.path.relpath(full_path, '.').replace('\\', '/')
            
            with open(full_path, 'r', encoding='utf-8', errors='ignore') as f_in:
                code_content = f_in.read()
            
            meta = meta_lookup.get(rel_path, {
                'file_name': file,
                'file_path': rel_path,
                'category': 'ملف مصدري (Source File)',
                'role': f'ملف برمجي ينتمي لمشروع تاسك فلو في المسار {rel_path}.',
                'key_items': 'كلاسات ودوال الواجهة الأمامية.',
                'doctor_q': 'ما هو دور هذا الملف في المنظومة؟',
                'doctor_a': f'يؤدي وظيفة تكاملية ضمن مجلد lib لدعم تشغيل الواجهة وإدارة البيانات.'
            })
            
            if file_count > 1:
                doc.add_page_break()
                
            add_h1(doc, f'[{file_count}/29] 📄 {meta["file_name"]}')
            
            # Info Box
            table_info = doc.add_table(rows=4, cols=1)
            table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_info.autofit = False
            
            c0 = table_info.rows[0].cells[0]
            c0.width = Inches(7.0)
            set_cell_background(c0, 'F1F5F9')
            set_cell_margins(c0, top=70, bottom=70, left=150, right=150)
            p_p = c0.paragraphs[0]
            p_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_p._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_rtl_text(p_p, 'المسار البرمجي: ', font_name='Arial', size_pt=9.5, bold=True, color_rgb=(71, 85, 105))
            add_rtl_text(p_p, rel_path, font_name='Consolas', size_pt=9.5, color_rgb=(15, 23, 42))

            c1 = table_info.rows[1].cells[0]
            c1.width = Inches(7.0)
            set_cell_background(c1, 'FFFFFF')
            set_cell_margins(c1, top=90, bottom=90, left=150, right=150)
            p_r = c1.paragraphs[0]
            p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_r._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_rtl_text(p_r, '📌 الشرح والدور الهندسي: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(30, 41, 59))
            add_rtl_text(p_r, meta['role'], font_name='Arial', size_pt=10, color_rgb=(51, 65, 85))

            c2 = table_info.rows[2].cells[0]
            c2.width = Inches(7.0)
            set_cell_background(c2, 'F8FAFC')
            set_cell_margins(c2, top=90, bottom=90, left=150, right=150)
            p_k = c2.paragraphs[0]
            p_k.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_k._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_rtl_text(p_k, '🔑 أهم المكونات والدوال: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(15, 118, 110))
            add_rtl_text(p_k, meta['key_items'], font_name='Arial', size_pt=10, color_rgb=(51, 65, 85))

            c3 = table_info.rows[3].cells[0]
            c3.width = Inches(7.0)
            set_cell_background(c3, 'EFF6FF')
            set_cell_margins(c3, top=100, bottom=100, left=150, right=150)
            p_q = c3.paragraphs[0]
            p_q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_q._p.get_or_add_pPr().append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
            add_rtl_text(p_q, '❓ سؤال الدكتور المتوقع: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(180, 83, 9))
            add_rtl_text(p_q, f'"{meta["doctor_q"]}"\n', font_name='Arial', size_pt=10, italic=True, bold=True, color_rgb=(146, 64, 14))
            add_rtl_text(p_q, '💡 الإجابة النموذجية: ', font_name='Arial', size_pt=10, bold=True, color_rgb=(21, 128, 61))
            add_rtl_text(p_q, meta['doctor_a'], font_name='Arial', size_pt=10, color_rgb=(22, 101, 52))
            
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(6)
            p_sp.paragraph_format.space_after = Pt(2)
            
            # Code block section
            add_h2(doc, f'💻 الكود المصدري للملف ({len(code_content.splitlines())} سطراً):')
            add_code_box(doc, code_content)
            
            print(f'Processed [{file_count}/29]: {rel_path}')

doc_out = 'docs/TaskFlow_Lib_Source_Code_Manual.docx'
doc.save(doc_out)
print(f'[3/4] Successfully saved Word document to {doc_out}')

print('[4/4] Converting to PDF...')
pdf_out = 'docs/TaskFlow_Lib_Source_Code_Manual.pdf'
docx2pdf.convert(doc_out, pdf_out)
print(f'SUCCESS: Complete Lib Code PDF generated at {pdf_out}')
